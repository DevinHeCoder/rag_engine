"""Word(.docx) 解析器：按一级标题分块；无标题时整个文档作为一个块。"""
from pathlib import Path
from typing import List, Optional, Tuple

from docx import Document as DocxDocument

from src.common.exceptions import DocumentParsingError
from src.document_parser.base import BaseDocumentParser, Document
from src.document_parser.registry import document_parser_registry


def _is_heading1(style_name: str) -> bool:
    """判断段落样式是否为一级标题（兼容中英文样式名）。"""
    s = (style_name or "").lower().strip()
    return s.startswith("heading 1") or s.startswith("标题 1")


@document_parser_registry.register("word")
class WordParser(BaseDocumentParser):
    """按 Heading 1 分块的 Word 解析器。"""

    def parse(self, file_path: str) -> List[Document]:
        p = Path(file_path)
        if not p.is_file():
            raise DocumentParsingError(f"Word 文件不存在: {p}")
        try:
            docx = DocxDocument(str(p))
        except Exception as e:
            raise DocumentParsingError(f"读取 Word 文件失败: {p}，原因: {e}") from e

        blocks = self._split_by_heading1(docx)
        docs: List[Document] = []
        for heading, paras in blocks:
            content = "\n".join(paras).strip()
            if not content and not heading:
                continue
            docs.append(Document(
                doc_id=f"{p.stem}::{len(docs):03d}",
                content=content,
                metadata={
                    "source": str(p),
                    "file_name": p.name,
                    "file_type": "docx",
                    "section": heading,
                    "index": len(docs),
                },
            ))
        return docs

    @staticmethod
    def _split_by_heading1(docx) -> List[Tuple[Optional[str], List[str]]]:
        """遍历段落，按 Heading 1 切分。返回 [(heading, [paragraph_texts]), ...]。"""
        blocks: List[Tuple[Optional[str], List[str]]] = []
        current_heading: Optional[str] = None
        current_paras: List[str] = []

        for para in docx.paragraphs:
            if _is_heading1(para.style.name):
                if current_paras or current_heading is not None:
                    blocks.append((current_heading, list(current_paras)))
                current_heading = para.text.strip() or None
                current_paras = []
            else:
                t = para.text.strip()
                if t:
                    current_paras.append(t)

        if current_paras or current_heading is not None:
            blocks.append((current_heading, list(current_paras)))
        return blocks
