"""P1 文档解析层单测：Markdown / PDF / Word 三种解析器 + 注册表 + dispatcher。"""
from pathlib import Path

import pytest

from src.common.exceptions import DocumentParsingError
from src.document_parser import (
    Document,
    document_parser_registry,
    get_parser_name,
    parse_file,
)
from src.document_parser.md_parser import MarkdownParser
from src.document_parser.pdf_parser import PdfParser
from src.document_parser.word_parser import WordParser


# ---------------- 辅助：生成样例文件 ----------------
def _make_md(path: Path) -> None:
    path.write_text(
        "# Title A\n\nContent A line 1.\nContent A line 2.\n\n"
        "# Title B\n\nContent B.\n",
        encoding="utf-8",
    )


def _make_docx(path: Path) -> None:
    from docx import Document
    doc = Document()
    doc.add_heading("Title One", level=1)
    doc.add_paragraph("Content under title one.")
    doc.add_heading("Title Two", level=1)
    doc.add_paragraph("Content under title two.")
    doc.save(str(path))


def _make_pdf(path: Path, text: str = "Hello PDF World") -> None:
    """生成一个最小的单页有效 PDF（手写字节，不依赖额外库）。"""
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
         b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>"),
        None,  # 占位，后面填 stream
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    stream = f"BT /F1 12 Tf 100 700 Td ({text}) Tj ET\n".encode("latin-1")
    objects[3] = (
        b"<< /Length " + str(len(stream)).encode()
        + b" >>\nstream\n" + stream + b"endstream"
    )

    pdf = b"%PDF-1.4\n"
    offsets = []
    for i, obj in enumerate(objects, start=1):
        offsets.append(len(pdf))
        pdf += f"{i} 0 obj\n".encode() + obj + b"\nendobj\n"

    xref_pos = len(pdf)
    pdf += b"xref\n"
    pdf += f"0 {len(objects) + 1}\n".encode()
    pdf += b"0000000000 65535 f \n"
    for off in offsets:
        pdf += f"{off:010d} 00000 n \n".encode()
    pdf += b"trailer\n"
    pdf += f"<< /Size {len(objects) + 1} /Root 1 0 R >>\n".encode()
    pdf += b"startxref\n"
    pdf += f"{xref_pos}\n".encode()
    pdf += b"%%EOF\n"
    path.write_bytes(pdf)


# ---------------- 注册表测试 ----------------
class TestRegistry:
    def test_all_parsers_registered(self):
        assert document_parser_registry.has("markdown")
        assert document_parser_registry.has("pdf")
        assert document_parser_registry.has("word")
        assert set(document_parser_registry.names) == {"markdown", "pdf", "word"}

    def test_create_by_name(self):
        assert isinstance(document_parser_registry.create("markdown"), MarkdownParser)
        assert isinstance(document_parser_registry.create("pdf"), PdfParser)
        assert isinstance(document_parser_registry.create("word"), WordParser)


# ---------------- Markdown 解析器 ----------------
class TestMarkdownParser:
    def test_split_by_h1(self, tmp_path):
        p = tmp_path / "doc.md"
        _make_md(p)
        docs = MarkdownParser().parse(str(p))
        assert len(docs) == 2
        assert all(isinstance(d, Document) for d in docs)
        assert docs[0].metadata["section"] == "Title A"
        assert docs[1].metadata["section"] == "Title B"
        assert "Content A line 1." in docs[0].content
        assert docs[0].metadata["file_type"] == "md"
        assert docs[0].metadata["file_name"] == "doc.md"
        assert docs[0].doc_id == "doc::000"

    def test_preamble_without_heading(self, tmp_path):
        p = tmp_path / "pre.md"
        p.write_text("Preamble text.\n\n# H1\nBody.\n", encoding="utf-8")
        docs = MarkdownParser().parse(str(p))
        assert len(docs) == 2
        assert docs[0].metadata["section"] is None
        assert "Preamble text." in docs[0].content

    def test_no_heading_single_block(self, tmp_path):
        p = tmp_path / "plain.md"
        p.write_text("Just text.\nNo heading.\n", encoding="utf-8")
        docs = MarkdownParser().parse(str(p))
        assert len(docs) == 1
        assert docs[0].metadata["section"] is None

    def test_missing_file_raises(self):
        with pytest.raises(DocumentParsingError):
            MarkdownParser().parse("/no/such/file.md")


# ---------------- PDF 解析器 ----------------
class TestPdfParser:
    def test_extract_text(self, tmp_path):
        p = tmp_path / "doc.pdf"
        _make_pdf(p, text="Hello PDF World")
        docs = PdfParser().parse(str(p))
        assert len(docs) >= 1
        assert docs[0].metadata["file_type"] == "pdf"
        assert docs[0].metadata["page"] == 1
        assert "Hello PDF World" in docs[0].content
        assert docs[0].doc_id.startswith("doc::page-")

    def test_missing_file_raises(self):
        with pytest.raises(DocumentParsingError):
            PdfParser().parse("/no/such/file.pdf")


# ---------------- Word 解析器 ----------------
class TestWordParser:
    def test_split_by_heading1(self, tmp_path):
        p = tmp_path / "doc.docx"
        _make_docx(p)
        docs = WordParser().parse(str(p))
        assert len(docs) == 2
        assert docs[0].metadata["section"] == "Title One"
        assert docs[1].metadata["section"] == "Title Two"
        assert "Content under title one." in docs[0].content
        assert docs[0].metadata["file_type"] == "docx"

    def test_no_heading_single_block(self, tmp_path):
        from docx import Document
        p = tmp_path / "plain.docx"
        doc = Document()
        doc.add_paragraph("Just a paragraph.")
        doc.save(str(p))
        docs = WordParser().parse(str(p))
        assert len(docs) == 1
        assert docs[0].metadata["section"] is None

    def test_missing_file_raises(self):
        with pytest.raises(DocumentParsingError):
            WordParser().parse("/no/such/file.docx")


# ---------------- dispatcher ----------------
class TestDispatcher:
    def test_get_parser_name(self):
        assert get_parser_name(".md") == "markdown"
        assert get_parser_name(".MD") == "markdown"
        assert get_parser_name(".pdf") == "pdf"
        assert get_parser_name(".docx") == "word"

    def test_unsupported_extension_raises(self):
        with pytest.raises(DocumentParsingError):
            get_parser_name(".txt")

    def test_parse_file_md(self, tmp_path):
        p = tmp_path / "doc.md"
        _make_md(p)
        docs = parse_file(str(p))
        assert len(docs) == 2
        assert docs[0].metadata["file_type"] == "md"

    def test_parse_file_pdf(self, tmp_path):
        p = tmp_path / "doc.pdf"
        _make_pdf(p)
        docs = parse_file(str(p))
        assert len(docs) >= 1

    def test_parse_file_docx(self, tmp_path):
        p = tmp_path / "doc.docx"
        _make_docx(p)
        docs = parse_file(str(p))
        assert len(docs) == 2

    def test_parse_file_missing_raises(self, tmp_path):
        with pytest.raises(DocumentParsingError):
            parse_file(str(tmp_path / "nope.md"))
